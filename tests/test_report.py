import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from kiosk_sales.__main__ import main
from kiosk_sales.loader import load_all, read_sales_file
from kiosk_sales.report import by_product, store_overview

HEADER = ["번호", "대분류", "상품코드", "상품명", "판매수", "결제 합계",
          "현금", "카드", "포인트", "외상", "기타", "할인 합계"]


def make_kiosk_file(path, venue, month, items):
    """키오스크 '상품별 매출' 내보내기와 같은 양식의 엑셀을 만든다.
    items: (대분류, 상품코드, 상품명, 판매수, 카드, 기타)"""
    wb = Workbook()
    ws = wb.active
    ws["B1"] = f"{venue}의 {month}-01 ~ {month}-30 상품별 매출"
    ws.merge_cells("B1:M1")
    ws.append([])
    ws.append([None] + HEADER)
    rows = [[i, c, code, name, q, card + etc, 0, card, 0, 0, etc, 0]
            for i, (c, code, name, q, card, etc) in enumerate(items, 1)]
    total = ["합계", None, None, None] + [sum(r[k] for r in rows) for k in range(4, 12)]
    ws.append([None] + total)
    for r in rows:
        ws.append([None] + r)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


SEP = [
    ("커피", "00006", "범표라떼", 10, 50000, 5000),
    ("커피", "00005", "ICE 아메리카노", 5, 20000, 0),
    ("커피제품&포장", "00082", "드립백 세트", 2, 20000, 0),
    ("커피제품", "00082", "드립백 세트", 1, 10000, 0),   # 월 중 대분류 변경
    ("할인", "00154", "원두 구매 시 2천원 할인", 1, -2000, 0),
]
AUG = [("커피", "00006", "범표라떼", 8, 40000, 0)]


@pytest.fixture
def raw_dir(tmp_path):
    raw = tmp_path / "raw"
    make_kiosk_file(raw / "별내점" / "2026-09.xlsx", "별내농업협동조합 하나로마트", "2026-09", SEP)
    make_kiosk_file(raw / "별내점" / "2026-08.xlsx", "별내농업협동조합 하나로마트", "2026-08", AUG)
    make_kiosk_file(raw / "강남점" / "2026-09.xlsx", "강남점", "2026-09", AUG)
    return raw


def test_read_sales_file(raw_dir):
    sf = read_sales_file(raw_dir / "별내점" / "2026-09.xlsx")
    assert sf.venue == "별내농업협동조합 하나로마트"
    assert sf.month == "2026-09"
    assert len(sf.items) == 5
    assert sf.items["amount"].sum() == 103000
    assert sf.items.loc[0, "code"] == "00006"


def test_total_mismatch_rejected(tmp_path):
    path = make_kiosk_file(tmp_path / "x.xlsx", "별내점", "2026-09", AUG)
    wb = load_workbook(path)
    wb.active["G4"] = 1  # 합계 행 결제 합계 조작
    wb.save(path)
    with pytest.raises(ValueError, match="합계"):
        read_sales_file(path)


def test_store_overview_and_products(raw_dir):
    df = load_all(raw_dir)
    s = df[df["store"] == "별내점"]
    cur, prev = s[s["month"] == "2026-09"], s[s["month"] == "2026-08"]

    ov = store_overview(cur, prev, royalty_rate=0.03).set_index("항목")["값"]
    assert ov["총매출(결제 합계)"] == 103000
    assert ov["전월 총매출"] == 40000
    assert ov["전월대비(%)"] == 157.5
    assert ov["로열티"] == 3090

    prod = by_product(cur, prev).set_index("상품코드")
    assert prod.loc["00082", "판매수량"] == 3          # 대분류가 달라도 상품코드로 합침
    assert prod.loc["00082", "대분류"] == "커피제품&포장"  # 판매수 많은 쪽이 대표
    assert prod.loc["00006", "전월대비(%)"] == 37.5


def test_cli_add_and_report(raw_dir, tmp_path):
    src = make_kiosk_file(tmp_path / "upload.xlsx", "별내농업협동조합 하나로마트", "2026-10", AUG)
    assert main(["--raw-dir", str(raw_dir), "add", str(src)]) == 0
    assert (raw_dir / "별내점" / "2026-10.xlsx").exists()
    assert main(["--raw-dir", str(raw_dir), "add", str(src)]) == 1   # 이미 있으면 거부

    out = tmp_path / "reports"
    assert main(["--raw-dir", str(raw_dir), "--out-dir", str(out), "report", "2026-09"]) == 0
    assert load_workbook(out / "별내점" / "별내점_2026-09_매출리포트.xlsx").sheetnames == \
        ["요약", "대분류별", "상품별", "월별 추이"]
    summary = pd.read_excel(out / "전체" / "2026-09_가맹점별요약.xlsx")
    assert list(summary["가맹점"]) == ["별내점", "강남점", "합계"]
    assert summary["총매출"].iloc[-1] == 143000


def test_unknown_store_requires_flag(tmp_path):
    src = make_kiosk_file(tmp_path / "u.xlsx", "모르는매장", "2026-09", AUG)
    assert main(["--raw-dir", str(tmp_path / "raw"), "add", str(src)]) == 1
    assert main(["--raw-dir", str(tmp_path / "raw"), "add", str(src), "--store", "신규점"]) == 0
